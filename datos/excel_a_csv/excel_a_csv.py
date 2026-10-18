# Convierte una hoja de un libro Excel (.xlsx) a un archivo CSV (UTF-8, coma) junto al original.
# Uso: python excel_a_csv.py archivo.xlsx [hoja]   (sin hoja, usa la hoja activa)
# Requiere: openpyxl (pip install -r requirements.txt). No lee el formato antiguo .xls.
# Autor: Eric Armijo
import csv
import os
import sys

from openpyxl import load_workbook


def excel_a_csv(archivo_excel, hoja=None):
    ruta, _ = os.path.splitext(archivo_excel)
    archivo_csv = ruta + ".csv"

    # data_only=True: escribe los valores calculados y no el texto de las fórmulas
    libro = load_workbook(filename=archivo_excel, data_only=True, read_only=True)
    if hoja is None:
        hoja_excel = libro.active
    elif hoja in libro.sheetnames:
        hoja_excel = libro[hoja]
    else:
        sys.exit(f"La hoja '{hoja}' no existe. Hojas disponibles: {', '.join(libro.sheetnames)}")

    with open(archivo_csv, "w", newline="", encoding="utf-8") as salida:
        escritor = csv.writer(salida, delimiter=",")
        for fila in hoja_excel.iter_rows(values_only=True):
            escritor.writerow(fila)
    libro.close()
    print(f"Escrito: {archivo_csv}")


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        print("Uso: python excel_a_csv.py <archivo.xlsx> [hoja]")
        sys.exit(1)
    excel_a_csv(sys.argv[1], sys.argv[2] if len(sys.argv) == 3 else None)
